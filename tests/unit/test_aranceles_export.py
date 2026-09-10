"""Export a Excel de OPERACIONES → ARANCELES (`api/services/aranceles_export.py`).

Sin base: se mockea `_q` y se relee el .xlsx generado con openpyxl para fijar el
CONTRATO del archivo — tres hojas, cabecera con el rango como FECHA (no texto),
importes/conteos como NÚMERO, claves como TEXTO, y que TODOS los filtros de la
vista lleguen al WHERE de las tres tablas.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest

from api.services import aranceles_export as svc
from api.services import operaciones_sql as ops


def _rows_for(sql: str) -> list[dict]:
    """Filas falsas según qué agrupa la query (mismo shape que devuelve Postgres)."""
    if "AS periodo" in sql:                # serie del gráfico (solo la vista)
        return [{"periodo": "2026-09-01", "ar": 2000.0}]
    if "AS clave" in sql:
        return [{"clave": "Bonos", "ar": 1500.5, "n": 3}, {"clave": "Acciones", "ar": 499.5, "n": 1}]
    if "GROUP BY id_cuenta" in sql:      # dim operador: se agrupa por cuenta
        return [{"id_cuenta": 805, "ar": 2000.0, "n": 4}]
    if "AS denominacion" in sql:
        return [{"denominacion": "MOLLO NICOLAS", "ar": 2000.0, "n": 4}]
    if "AS instrumento" in sql:
        return [{"instrumento": "0123", "ar": 2000.0, "n": 4}]   # ticker numérico → TEXTO
    raise AssertionError(f"query inesperada: {sql}")


@pytest.fixture
def mock_q(monkeypatch):
    llamadas: list[tuple[str, dict]] = []

    def _q(sql, params=None):
        llamadas.append((sql, dict(params or {})))
        return _rows_for(sql)

    monkeypatch.setattr(ops, "_q", _q)
    monkeypatch.setattr(ops, "_operador_map", lambda: {"805": "Juan"})
    return llamadas


def test_tres_hojas_con_tipos_correctos(mock_q):
    openpyxl = pytest.importorskip("openpyxl")
    contenido, nombre = svc.export_xlsx(moneda="ARS", desde="2026-09-01", hasta="2026-09-10")
    assert nombre == "aranceles_20260901_20260910.xlsx"
    wb = openpyxl.load_workbook(BytesIO(contenido))
    assert wb.sheetnames == ["Consolidado", "Por cliente", "Por instrumento"]

    ws = wb["Consolidado"]
    # Rango elegido como celdas de FECHA con formato de fecha (no strings).
    assert ws["A2"].value == "Desde" and isinstance(ws["B2"].value, datetime)
    assert ws["B2"].value.date().isoformat() == "2026-09-01"
    assert ws["A3"].value == "Hasta" and ws["B3"].value.date().isoformat() == "2026-09-10"
    assert ws["B2"].number_format == "dd/mm/yyyy"
    assert ws["B4"].value == "ARS"
    # Las CUATRO dimensiones del selector, en el orden de la vista.
    titulos = [c.value for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("POR ")]
    assert titulos == ["POR NIVEL 3", "POR OPERACIÓN", "POR MERCADO", "POR OPERADOR"]
    # Primer bloque: header + filas tipadas + TOTAL.
    fila_hdr = next(c.row for c in ws["A"] if c.value == "POR NIVEL 3") + 1
    assert [c.value for c in ws[fila_hdr]][:4] == ["Nivel 3", "Aranceles (ARS)", "Boletos", "% del total"]
    r = fila_hdr + 1
    assert ws.cell(r, 1).value == "Bonos" and ws.cell(r, 1).number_format == "@"
    assert ws.cell(r, 2).value == 1500.5 and isinstance(ws.cell(r, 2).value, float)
    assert ws.cell(r, 2).number_format == "#,##0.00"
    assert ws.cell(r, 3).value == 3 and isinstance(ws.cell(r, 3).value, int)
    assert ws.cell(r, 4).value == pytest.approx(1500.5 / 2000.0)
    assert ws.cell(r, 4).number_format == "0.0%"
    assert ws.cell(r + 2, 1).value == "TOTAL" and ws.cell(r + 2, 2).value == 2000.0
    # Bloque OPERADOR resuelto por el mapa cuenta→operador.
    fila_op = next(c.row for c in ws["A"] if c.value == "POR OPERADOR") + 2
    assert ws.cell(fila_op, 1).value == "Juan" and ws.cell(fila_op, 3).value == 4

    # Por cliente / por instrumento: la clave queda TEXTO aunque sea numérica.
    wc = wb["Por cliente"]
    fila = next(c.row for c in wc["B"] if c.value == "Aranceles (ARS)") + 1
    assert wc.cell(fila - 1, 1).value == "Cliente"
    assert wc.cell(fila, 1).value == "MOLLO NICOLAS" and wc.cell(fila, 2).value == 2000.0
    wi = wb["Por instrumento"]
    fila = next(c.row for c in wi["B"] if c.value == "Aranceles (ARS)") + 1
    assert wi.cell(fila - 1, 1).value == "Instrumento"
    assert wi.cell(fila, 1).value == "0123" and isinstance(wi.cell(fila, 1).value, str)
    assert wi.cell(fila, 1).number_format == "@"


def test_filtros_llegan_a_todas_las_tablas(mock_q):
    pytest.importorskip("openpyxl")
    contenido, _ = svc.export_xlsx(
        moneda="USD", desde="2026-09-10", hasta="2026-09-01",   # invertido → se normaliza
        cuenta="MOLLO NICOLAS", instrumento="AL30", sel_dim="Bonos", dim="nivel3",
        segmento="Minorista", operador="juan@aca.com", scope=("805",),
    )
    # 4 bloques + cliente + instrumento = 6 queries, TODAS con todos los filtros.
    assert len(mock_q) == 6
    for sql, p in mock_q:
        assert "denominacion = %(f_cuenta)s" in sql and p["f_cuenta"] == "MOLLO NICOLAS"
        assert "instrumento = %(f_instr)s" in sql and p["f_instr"] == "AL30"
        assert "nivel_3 = %(f_dim)s" in sql and p["f_dim"] == "Bonos"
        assert "segmento = ANY(%(segmento)s)" in sql and p["segmento"] == ["Minorista"]
        assert "operador_email = %(f_op)s" in sql and p["f_op"] == "juan@aca.com"
        assert "id_cuenta = ANY(%(scope)s)" in sql and p["scope"] == ["805"]
        assert p["desde"] == "2026-09-01" and p["hasta"] == "2026-09-10"
        assert "ABS(arancel) /" in sql   # USD: arancel dolarizado por boleto
    # La cabecera de cada hoja deja constancia de los filtros.
    import openpyxl
    ws = openpyxl.load_workbook(BytesIO(contenido))["Por cliente"]
    labels = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, 11)}
    assert labels["Segmento"] == "Minorista" and labels["Operador"] == "juan@aca.com"
    assert labels["Selección NIVEL 3"] == "Bonos" and labels["Cliente"] == "MOLLO NICOLAS"
    assert labels["Instrumento"] == "AL30" and labels["Moneda"] == "USD"


def test_sel_dim_operador_usa_subquery(mock_q):
    pytest.importorskip("openpyxl")
    svc.export_xlsx(moneda="ARS", desde="2026-09-01", hasta="2026-09-10",
                    sel_dim="Juan", dim="operador")
    for sql, p in mock_q:
        assert "COALESCE(o.nombre, o.email) = %(sel_op)s" in sql and p["sel_op"] == "Juan"


def test_fecha_invalida_es_value_error(mock_q):
    with pytest.raises(ValueError, match="desde inválida"):
        svc.export_xlsx(moneda="ARS", desde="10/09/2026", hasta="2026-09-10")
    assert mock_q == []   # se valida ANTES de consultar


def test_ops_aranceles_sigue_igual_tras_el_refactor(mock_q, monkeypatch):
    """La vista no cambió: por_dim ignora sel_dim, por_cuenta ignora cuenta, etc."""
    out = ops.ops_aranceles(moneda="ARS", desde="2026-09-01", hasta="2026-09-10", agg="DIARIO",
                            cuenta="X", instrumento="Y", sel_dim="Bonos", dim="nivel3")
    assert out["total"] == 2000.0 and out["por_dim"][0]["clave"] == "Bonos"
    por_sql = {s for s, _ in mock_q}
    dim_sql = next(s for s in por_sql if "AS clave" in s)
    assert "f_dim" not in dim_sql and "f_cuenta" in dim_sql and "f_instr" in dim_sql
    cta_sql = next(s for s in por_sql if "AS denominacion" in s)
    assert "f_cuenta" not in cta_sql and "f_dim" in cta_sql and "f_instr" in cta_sql
    ins_sql = next(s for s in por_sql if "AS instrumento" in s)
    assert "f_instr" not in ins_sql and "f_dim" in ins_sql and "f_cuenta" in ins_sql
