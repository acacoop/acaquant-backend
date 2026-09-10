"""api/services/aranceles_export.py — export a Excel de OPERACIONES → ARANCELES.

Servicio PURO (sin FastAPI). Arma un .xlsx de TRES hojas con lo que la vista
muestra para el rango y los filtros elegidos:

  * CONSOLIDADO      — la tabla izquierda superior en TODAS sus dimensiones
                       (NIVEL 3 · OPERACIÓN · MERCADO · OPERADOR), un bloque
                       por dimensión con su fila TOTAL. En pantalla se ve una
                       dimensión por vez; el archivo las trae todas.
  * POR CLIENTE      — la tabla derecha superior.
  * POR INSTRUMENTO  — la tabla derecha inferior.

Cada hoja abre con la cabecera del corte: DESDE / HASTA como celdas de FECHA
(no texto), moneda y los filtros activos. Tipado de celdas: los importes y
conteos van como NÚMERO con formato (`#,##0.00` / `#,##0`), el % como fracción
con formato `0.0%`, las claves como TEXTO (`@`, así un ticker numérico no se
vuelve número). Los datos salen de `operaciones_sql.ops_aranceles_export`, que
aplica TODOS los filtros a las tres tablas (a diferencia de la pantalla, donde
cada tabla ignora su propia selección para poder des-seleccionar).

openpyxl con import lazy para no tumbar la API si falta la lib (REGLA #1),
mismo patrón que `senebis._armar_xlsx`.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from api.services import operaciones_sql as _ops_sql

_TZ_AR = ZoneInfo("America/Argentina/Buenos_Aires")

_FMT_FECHA = "dd/mm/yyyy"
_FMT_FECHA_HORA = "dd/mm/yyyy hh:mm"
_FMT_IMPORTE = "#,##0.00"
_FMT_ENTERO = "#,##0"
_FMT_PCT = "0.0%"
_FMT_TEXTO = "@"

# Etiquetas de la cabecera de filtros (en el orden en que aparecen en la vista).
_DIM_LABEL = dict(_ops_sql._ARANCEL_DIMS)


def _parse_fecha(s: str, nombre: str) -> date:
    try:
        return date.fromisoformat(str(s or "").strip())
    except ValueError as e:
        raise ValueError(f"{nombre} inválida: {s!r} (esperado YYYY-MM-DD)") from e


def _estilos():
    from openpyxl.styles import Alignment, Font, PatternFill
    return {
        "titulo": Font(bold=True, size=13, color="1F4E79"),
        "label": Font(bold=True),
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill("solid", fgColor="1F4E79"),
        "bloque": Font(bold=True, size=11, color="1F4E79"),
        "total": Font(bold=True),
        "derecha": Alignment(horizontal="right"),
    }


def _cabecera(ws, st: dict, titulo: str, data: dict, desde: date, hasta: date) -> int:
    """Escribe título + corte + filtros arriba de la hoja. Devuelve la próxima
    fila libre (después de una fila en blanco)."""
    f = data["filtros"]
    ws.cell(row=1, column=1, value=titulo).font = st["titulo"]
    filas: list[tuple[str, object, str | None]] = [
        ("Desde", desde, _FMT_FECHA),
        ("Hasta", hasta, _FMT_FECHA),
        ("Moneda", data["moneda"], _FMT_TEXTO),
        ("Segmento", f.get("segmento") or "Todos", _FMT_TEXTO),
        ("Operador", f.get("operador") or "Todos", _FMT_TEXTO),
        (f"Selección {_DIM_LABEL.get(data['dim'], data['dim'])}",
         f.get("sel_dim") or "Todos", _FMT_TEXTO),
        ("Cliente", f.get("cuenta") or "Todos", _FMT_TEXTO),
        ("Instrumento", f.get("instrumento") or "Todos", _FMT_TEXTO),
        ("Generado", datetime.now(_TZ_AR).replace(tzinfo=None), _FMT_FECHA_HORA),
    ]
    for i, (label, valor, fmt) in enumerate(filas, start=2):
        ws.cell(row=i, column=1, value=label).font = st["label"]
        c = ws.cell(row=i, column=2, value=valor)
        if fmt:
            c.number_format = fmt
    return len(filas) + 3


def _tabla(ws, st: dict, fila: int, clave_hdr: str, moneda: str, filas: list[dict],
           key: str, *, titulo: str | None = None) -> int:
    """Escribe una tabla (opcionalmente con título de bloque): header azul, filas
    tipadas (texto / importe / entero / %), y fila TOTAL. Devuelve la próxima fila
    libre (después de una fila en blanco)."""
    if titulo:
        ws.cell(row=fila, column=1, value=titulo).font = st["bloque"]
        fila += 1
    headers = (clave_hdr, f"Aranceles ({moneda})", "Boletos", "% del total")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=fila, column=col, value=h)
        c.font = st["header_font"]
        c.fill = st["header_fill"]
        if col > 1:
            c.alignment = st["derecha"]
    fila += 1
    total_ar = sum(float(r["arancel"]) for r in filas)
    total_n = sum(int(r["n"]) for r in filas)
    for r in filas:
        c = ws.cell(row=fila, column=1, value=str(r[key]))
        c.number_format = _FMT_TEXTO
        ws.cell(row=fila, column=2, value=float(r["arancel"])).number_format = _FMT_IMPORTE
        ws.cell(row=fila, column=3, value=int(r["n"])).number_format = _FMT_ENTERO
        pct = (float(r["arancel"]) / total_ar) if total_ar else 0.0
        ws.cell(row=fila, column=4, value=pct).number_format = _FMT_PCT
        fila += 1
    if not filas:
        ws.cell(row=fila, column=1, value="sin datos")
        fila += 1
    ws.cell(row=fila, column=1, value="TOTAL").font = st["total"]
    c = ws.cell(row=fila, column=2, value=round(total_ar, 2))
    c.number_format = _FMT_IMPORTE
    c.font = st["total"]
    c = ws.cell(row=fila, column=3, value=total_n)
    c.number_format = _FMT_ENTERO
    c.font = st["total"]
    c = ws.cell(row=fila, column=4, value=1.0 if filas else 0.0)
    c.number_format = _FMT_PCT
    c.font = st["total"]
    return fila + 2


def _anchos(ws, primera: int = 34) -> None:
    for letra, ancho in (("A", primera), ("B", 20), ("C", 12), ("D", 12)):
        ws.column_dimensions[letra].width = ancho


def armar_xlsx(data: dict) -> bytes:
    """dict de `ops_aranceles_export` → bytes del .xlsx (3 hojas)."""
    try:
        from openpyxl import Workbook
    except ImportError as e:  # lazy: sin openpyxl la API sigue viva (REGLA #1)
        raise RuntimeError(
            "openpyxl no está instalado — correr `pip install -r requirements.txt` "
            "en el venv del Droplet") from e
    desde = _parse_fecha(data["desde"], "desde")
    hasta = _parse_fecha(data["hasta"], "hasta")
    moneda = data["moneda"]
    st = _estilos()
    wb = Workbook()

    # 1) CONSOLIDADO: un bloque por dimensión del selector.
    ws = wb.active
    ws.title = "Consolidado"
    fila = _cabecera(ws, st, "ARANCELES — CONSOLIDADO", data, desde, hasta)
    for b in data["bloques"]:
        fila = _tabla(ws, st, fila, b["titulo"].capitalize(), moneda, b["filas"], "clave",
                      titulo=f"POR {b['titulo']}")
    _anchos(ws)

    # 2) POR CLIENTE.
    ws = wb.create_sheet("Por cliente")
    fila = _cabecera(ws, st, "ARANCELES — POR CLIENTE", data, desde, hasta)
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)
    _tabla(ws, st, fila, "Cliente", moneda, data["por_cuenta"], "denominacion")
    _anchos(ws, primera=44)

    # 3) POR INSTRUMENTO.
    ws = wb.create_sheet("Por instrumento")
    fila = _cabecera(ws, st, "ARANCELES — POR INSTRUMENTO", data, desde, hasta)
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)
    _tabla(ws, st, fila, "Instrumento", moneda, data["por_instrumento"], "instrumento")
    _anchos(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_xlsx(
    *, moneda: str, desde: str, hasta: str, cuenta: str | None = None,
    instrumento: str | None = None, sel_dim: str | None = None, segmento: str | None = None,
    dim: str = "nivel3", scope: tuple[str, ...] | None = None, operador: str | None = None,
) -> tuple[bytes, str]:
    """Devuelve (bytes del .xlsx, nombre de archivo). Valida las fechas ANTES de
    consultar (un rango inválido es 400, no una query rota)."""
    d, h = _parse_fecha(desde, "desde"), _parse_fecha(hasta, "hasta")
    if d > h:
        d, h = h, d   # misma normalización que hace la vista
    data = _ops_sql.ops_aranceles_export(
        moneda=moneda, desde=d.isoformat(), hasta=h.isoformat(), cuenta=cuenta,
        instrumento=instrumento, sel_dim=sel_dim, segmento=segmento, dim=dim,
        scope=scope, operador=operador,
    )
    nombre = f"aranceles_{d.strftime('%Y%m%d')}_{h.strftime('%Y%m%d')}.xlsx"
    return armar_xlsx(data), nombre
